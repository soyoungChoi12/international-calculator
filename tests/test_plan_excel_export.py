"""직원 해외출장 계획 Excel — 템플릿 사본에 계산·계획안을 넣는다."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from services.plan_excel_export import (
    TEMPLATE_PATH,
    build_plan_excel_bytes,
    build_party_plan_excel_bytes,
    plan_excel_filename,
)
from services.plan_parser import PlanDocument, parse_plan_text
from services.travel_calculator import PartyMember, PartyResult, StayInput, TravelInput, calculate_travel
from tests.test_plan_parser import SAMPLE_PLAN


def _kyoto_input() -> TravelInput:
    return TravelInput(
        role="팀장 및 팀원",
        grade="나",
        departure_date=date(2026, 6, 30),
        return_date=date(2026, 7, 3),
        lodging_nights=3,
        exchange_rate=1535,
        airfare_krw=520_300,
        lodging_actual_krw=0,
        preparation_krw=150_000,
        stays=[StayInput("일본", "교토", 3, "나", stay_days=4)],
    )


def _workbook(result, **kwargs):
    data = build_plan_excel_bytes(result, **kwargs)
    return load_workbook(BytesIO(data)), data


def test_plan_excel_filename_uses_year_name_and_date():
    assert plan_excel_filename("이한주", date(2026, 6, 22)) == "26년 직원 해외출장 계획_이한주(260622).xlsx"
    assert plan_excel_filename("", date(2026, 9, 4)) == "26년 직원 해외출장 계획_미기재(260904).xlsx"


def test_template_is_not_modified_by_plan_excel_export():
    before = TEMPLATE_PATH.read_bytes()
    result = calculate_travel(_kyoto_input())
    build_plan_excel_bytes(result, "이한주", title="전임", team="글로벌전략협업팀")
    assert TEMPLATE_PATH.read_bytes() == before


def test_kyoto_plan_excel_fills_identity_place_costs_and_fx():
    plan = parse_plan_text(SAMPLE_PLAN)
    plan = PlanDocument(
        **{
            **plan.__dict__,
            "domestic_krw": 30_000,
            "budget_category": "2026년 글로벌기업 협업 프로그램(AroundX) 사업 – 주관기관 운영 - 운영비 - 여비",
            "purpose_lines": ("2026 AroundX 정글 참여기업 5개사 일본 GTM 지원",),
            "airfare_note": " (김포 - 간사이공항)",
        }
    )
    result = calculate_travel(_kyoto_input())
    wb, data = _workbook(
        result,
        traveler_name="이한주",
        title="전임",
        team="글로벌전략협업팀",
        plan=plan,
        departure=date(2026, 6, 30),
        return_on=date(2026, 7, 3),
        approval_date=date(2026, 6, 22),
    )
    assert data[:2] == b"PK"
    ws = wb[wb.sheetnames[0]]
    assert ws["B1"].value == "2026 경기창조경제혁신센터 해외출장 계획"
    assert ws["F2"].value == "나 지역"
    assert ws["B6"].value == "전임"
    assert ws["C6"].value == "글로벌전략협업팀"
    assert ws["D6"].value == "이한주"
    assert ws["E6"].value == "2026.06.30~2026.07.03"
    assert ws["F6"].value == "일본, 교토(나)"
    assert "AroundX" in str(ws["G6"].value)
    assert "주관기관" in str(ws["H6"].value)
    assert ws["I6"].value == "=SUM(J6:O6)"
    assert ws["J6"].value == 520_300
    assert ws["K6"].value == 30_000
    assert ws["L6"].value == 566_415
    assert ws["M6"].value == 159_640
    assert ws["N6"].value == 300_860
    assert ws["O6"].value == 150_000
    assert "김포" in str(ws["J5"].value)
    assert ws["I9"].value == "환율"
    assert ws["J9"].value == 1535
    assert ws["J10"].value == "(26.06.22)"
    wb.close()


def test_missing_costs_are_zero():
    result = calculate_travel(
        TravelInput(
            role="팀장 및 팀원",
            grade="나",
            departure_date=date(2026, 6, 30),
            return_date=date(2026, 7, 3),
            lodging_nights=3,
            exchange_rate=1535,
            airfare_krw=0,
            lodging_actual_krw=0,
            preparation_krw=0,
            stays=[StayInput("일본", "교토", 3, "나", stay_days=4)],
        )
    )
    wb, _ = _workbook(
        result,
        traveler_name="이한주",
        title="전임",
        departure=date(2026, 6, 30),
        return_on=date(2026, 7, 3),
        approval_date=date(2026, 6, 22),
    )
    ws = wb[wb.sheetnames[0]]
    assert ws["J6"].value == 0
    assert ws["K6"].value == 0
    assert ws["O6"].value == 0
    assert ws["L6"].value == result.lodging.ceiling_krw
    wb.close()


def test_party_plan_excel_writes_one_sheet_per_person():
    staff = calculate_travel(_kyoto_input())
    head = calculate_travel(
        TravelInput(
            role="본부장",
            grade="나",
            departure_date=date(2026, 6, 30),
            return_date=date(2026, 7, 3),
            lodging_nights=3,
            exchange_rate=1535,
            airfare_krw=520_300,
            lodging_actual_krw=0,
            preparation_krw=150_000,
            stays=[StayInput("일본", "교토", 3, "나", stay_days=4)],
        )
    )
    party = [
        PartyResult(
            member=PartyMember(name="이한주", role="팀장 및 팀원", title="전임", team="글로벌전략협업팀"),
            result=staff,
        ),
        PartyResult(
            member=PartyMember(name="이종휘", role="본부장", title="본부장", team="글로벌전략협업팀"),
            result=head,
        ),
    ]
    data = build_party_plan_excel_bytes(
        party,
        departure=date(2026, 6, 30),
        return_on=date(2026, 7, 3),
        approval_date=date(2026, 6, 22),
    )
    wb = load_workbook(BytesIO(data))
    assert "이한주" in wb.sheetnames
    assert "이종휘" in wb.sheetnames
    assert wb["이한주"]["D6"].value == "이한주"
    assert wb["이한주"]["B6"].value == "전임"
    assert wb["이종휘"]["D6"].value == "이종휘"
    assert wb["이종휘"]["B6"].value == "본부장"
    assert wb["이한주"]["M6"].value == staff.daily.amount_krw
    assert wb["이종휘"]["M6"].value == head.daily.amount_krw
    wb.close()
