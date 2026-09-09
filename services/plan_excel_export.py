"""직원 해외출장 계획 Excel 생성.

원본 템플릿은 복사만 하고 수정하지 않는다.
"""

from __future__ import annotations

import re
import shutil
import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from services.plan_parser import PlanDocument
from services.travel_calculator import TravelResult

TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "26년 직원 해외출장 계획.xlsx"
SHEET_NAME = "2026"
TITLE_ORG = "경기창조경제혁신센터 해외출장 계획"


def plan_excel_filename(traveler_name: str, when: date) -> str:
    safe = re.sub(r'[\\/:*?"<>|]', "", (traveler_name or "").strip()) or "미기재"
    return f"{when.year % 100}년 직원 해외출장 계획_{safe}({when.strftime('%y%m%d')}).xlsx"


def build_plan_excel_bytes(
    result: TravelResult,
    traveler_name: str,
    *,
    title: str = "",
    team: str = "",
    plan: PlanDocument | None = None,
    departure: date | None = None,
    return_on: date | None = None,
    approval_date: date | None = None,
) -> bytes:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"해외출장 계획 템플릿을 찾을 수 없습니다: {TEMPLATE_PATH}")
    when = approval_date or date.today()
    start = departure or (plan.departure if plan else None)
    end = return_on or (plan.return_on if plan else None)
    with tempfile.TemporaryDirectory() as tmp:
        work_path = Path(tmp) / TEMPLATE_PATH.name
        shutil.copy2(TEMPLATE_PATH, work_path)
        wb = load_workbook(work_path)
        _fill_workbook(
            wb,
            result,
            traveler_name,
            title=title,
            team=team,
            plan=plan,
            start=start,
            end=end,
            approval_date=when,
        )
        buffer = BytesIO()
        wb.save(buffer)
        wb.close()
        return buffer.getvalue()


def _fill_workbook(
    wb,
    result: TravelResult,
    traveler_name: str,
    *,
    title: str,
    team: str,
    plan: PlanDocument | None,
    start: date | None,
    end: date | None,
    approval_date: date,
) -> None:
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
    year = (start or approval_date).year
    ws["B1"] = f"{year} {TITLE_ORG}"
    grade = _place_grade(result, plan)
    ws["F2"] = f"{grade} 지역" if grade else "지역"
    ws["B6"] = (title or "").strip()
    use_team = (team or "").strip()
    if not use_team:
        person = plan.travelers[0] if plan and getattr(plan, "travelers", None) else None
        use_team = (person.team if person else "") or ""
    ws["C6"] = use_team
    ws["D6"] = (traveler_name or "").strip()
    ws["E6"] = _period_text(start, end)
    ws["F6"] = _place_text(result, plan)
    ws["G6"] = _purpose_text(plan)
    ws["H6"] = _category_text(plan)
    ws["I6"] = "=SUM(J6:O6)"
    airfare, domestic, lodging, daily, meal, prep = _cost_row(result, plan)
    ws["J6"] = airfare
    ws["K6"] = domestic
    ws["L6"] = lodging
    ws["M6"] = daily
    ws["N6"] = meal
    ws["O6"] = prep
    note = _airfare_header(plan)
    if note:
        ws["J5"] = note
    ws["I9"] = "환율"
    ws["J9"] = _rate_value(result.exchange_rate)
    ws["J10"] = f"({approval_date.strftime('%y.%m.%d')})"


def _plan_get(plan: PlanDocument | None, name: str, default=""):
    if plan is None:
        return default
    return getattr(plan, name, default)


def _period_text(start: date | None, end: date | None) -> str:
    if start is None or end is None:
        return ""
    return f"{start.strftime('%Y.%m.%d')}~{end.strftime('%Y.%m.%d')}"


def _place_grade(result: TravelResult, plan: PlanDocument | None) -> str:
    stay = result.stays[0] if result.stays else None
    if stay and stay.grade:
        return stay.grade
    grade = (result.grade or "").split("/")[0].strip()
    if grade in {"가", "나", "다", "라"}:
        return grade
    plan_grade = _plan_get(plan, "grade") or ""
    return plan_grade if plan_grade in {"가", "나", "다", "라"} else ""


def _place_text(result: TravelResult, plan: PlanDocument | None) -> str:
    stay = result.stays[0] if result.stays else None
    country = (stay.country if stay else "") or _plan_get(plan, "country") or ""
    city = (stay.city if stay else "") or _plan_get(plan, "city") or ""
    grade = _place_grade(result, plan)
    place = ", ".join(part for part in (country.strip(), city.strip()) if part)
    if place and grade:
        return f"{place}({grade})"
    return place


def _purpose_text(plan: PlanDocument | None) -> str:
    if not plan:
        return ""
    lines = [str(line).strip().lstrip("◦").strip() for line in _plan_get(plan, "purpose_lines", ()) if str(line).strip()]
    if lines:
        return lines[0]
    return (plan.purpose or "").strip()


def _category_text(plan: PlanDocument | None) -> str:
    if not plan:
        return ""
    return (plan.budget_category or "").strip()


def _cost_row(result: TravelResult, plan: PlanDocument | None) -> tuple[int, int, int, int, int, int]:
    airfare = int(result.airfare_krw or 0)
    domestic = int(_plan_get(plan, "domestic_krw", 0) or 0)
    lodging = int(result.lodging.ceiling_krw or 0)
    daily = int(result.daily.amount_krw or 0)
    meal = int(result.meal.amount_krw or 0)
    prep = int(result.preparation_krw or 0)
    return airfare, domestic, lodging, daily, meal, prep


def _airfare_header(plan: PlanDocument | None) -> str:
    note = (_plan_get(plan, "airfare_note") or "").strip()
    if not note:
        return ""
    route = note.strip(" ()")
    if not route:
        return ""
    if "왕복" in route:
        return f"국제 항공료({route})"
    return f"국제 항공료({route} 왕복)"


def _rate_value(rate: float) -> int | float:
    if float(rate).is_integer():
        return int(rate)
    return rate
